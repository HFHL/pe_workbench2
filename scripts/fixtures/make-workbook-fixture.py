#!/usr/bin/env python3
"""Generate a small deterministic workbook used by scripts/test-workbook.mjs."""

import sys
from pathlib import Path

import openpyxl


def main() -> None:
    out = Path(sys.argv[1])
    revenue = int(sys.argv[2]) if len(sys.argv) > 2 else 100
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    model = wb.active
    model.title = "Model"
    model.append(["Metric", 2024, 2025, 2026])
    model.append(["Revenue", revenue, revenue + 10, revenue + 21])
    model.append(["COGS", 40, 44, 48])
    model.append(["Gross profit", "=B2-B3", "=C2-C3", "=D2-D3"])
    model.append(["Gross margin", "=B4/B2", "=C4/C2", "=D4/D2"])
    model.append(["Check total", "=SUM(B2:B5)", "=SUM(C2:C5)", "=SUM(D2:D5)"])
    model.append(["WACC", "=Assumption!B2", "=Assumption!B2", "=Assumption!B2"])
    model.append(["Tax rate", "=Assumption!B1", None, None])

    assumptions = wb.create_sheet("Assumption")
    assumptions.append(["Tax rate", 0.25])
    assumptions.append(["WACC", 0.1])

    chinese = wb.create_sheet("模型")
    chinese.append(["项目", 2025])
    chinese.append(["营业收入", 200])
    chinese.append(["营业成本", 80])
    chinese.append(["毛利", "=B2-B3"])
    chinese.append(["毛利率", "=B4/B2"])
    chinese.append(["收入增长率", "=B2/100-1"])
    chinese.append(["自由现金流", 50])
    chinese.append(["折现率", 0.1])
    chinese.append(["估值", "=B7/(1+B8)"])

    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
