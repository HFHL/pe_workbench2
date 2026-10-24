#!/usr/bin/env python3
"""Build the model-trace usage guide as a styled Word document."""

from __future__ import annotations

import re
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
SOURCE = DOCS / "model-trace-guide.md"
OUTPUT = DOCS / "模型追踪-功能实现与使用说明.docx"

INK = "1B1B19"
BODY = "262624"
MUTED = "6B675E"
FAINT = "8A857A"
GOLD = "8A6A2F"
LINE = "D8D2C3"
PAPER = "F6F2E8"
CARD = "FAF8F3"
CODE_BG = "F7F4ED"
LINK = "2F5D8A"
CONTENT_WIDTH_CM = 16.6

LATIN = "Helvetica Neue"
EAST_ASIA = "Hiragino Sans GB"
MONO = "Menlo"


def rgb(hex_value: str) -> RGBColor:
    value = hex_value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def set_run_font(run, *, name=LATIN, size=10.5, bold=False, italic=False, color=BODY):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = rgb(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), EAST_ASIA)


def shade_run(run, fill):
    rPr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    rPr.append(shd)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_margins(cell, top=60, bottom=60, left=90, right=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for key, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        element = OxmlElement(f"w:{key}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        tcMar.append(element)
    tcPr.append(tcMar)


def set_table_borders(table, color=LINE, size=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    existing = tblPr.find(qn("w:tblBorders"))
    if existing is not None:
        tblPr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    tblPr.append(borders)


def set_table_fixed(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    for tag in ("w:tblLayout", "w:tblW"):
        existing = tblPr.find(qn(tag))
        if existing is not None:
            tblPr.remove(existing)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), str(int(CONTENT_WIDTH_CM * 567)))
    width.set(qn("w:type"), "dxa")
    tblPr.append(width)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    trPr.append(element)


def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:cantSplit")
    element.set(qn("w:val"), "true")
    trPr.append(element)


def set_paragraph_border(paragraph, edge="bottom", color=LINE, size=6, space=4):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), str(space))
    element.set(qn("w:color"), color)
    pBdr.append(element)
    pPr.append(pBdr)


def add_page_field(paragraph, size=8, color=FAINT):
    run = paragraph.add_run()
    set_run_font(run, size=size, color=color)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_hyperlink(paragraph, text, url, size=10.5, bold=False, color=LINK):
    if not url:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=bold, color=BODY)
        return
    try:
        r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    except Exception:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=bold, color=BODY)
        return
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run_element = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), LATIN)
    rFonts.set(qn("w:hAnsi"), LATIN)
    rFonts.set(qn("w:eastAsia"), EAST_ASIA)
    rPr.append(rFonts)
    color_element = OxmlElement("w:color")
    color_element.set(qn("w:val"), color)
    rPr.append(color_element)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)
    size_element = OxmlElement("w:sz")
    size_element.set(qn("w:val"), str(int(size * 2)))
    rPr.append(size_element)
    if bold:
        bold_element = OxmlElement("w:b")
        rPr.append(bold_element)
    run_element.append(rPr)
    text_element = OxmlElement("w:t")
    text_element.set(qn("xml:space"), "preserve")
    text_element.text = text
    run_element.append(text_element)
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)


TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def resolve_target(target: str) -> str:
    if target.startswith(("http://", "https://", "mailto:")):
        return target
    if target.startswith("#"):
        return ""
    clean = target.split(" ", 1)[0].strip()
    path = (DOCS / clean).resolve()
    if path.exists():
        return path.as_uri()
    return ""


def add_inline(paragraph, text, *, size=10.5, bold=False, color=BODY, italic=False):
    position = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size=size, bold=bold, italic=italic, color=color)
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            add_inline(paragraph, token[2:-2], size=size, bold=True, color=color, italic=italic)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name=MONO, size=size - 1, bold=bold, italic=False, color="7A3E2E")
            shade_run(run, CODE_BG)
        else:
            link_match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link_match:
                add_hyperlink(paragraph, link_match.group(1), resolve_target(link_match.group(2)), size=size, bold=bold)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size, bold=bold, italic=italic, color=color)


def clean_md(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.strip()


def parse_blocks(source: str):
    lines = source.splitlines()
    blocks = []
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            language = line[3:].strip()
            index += 1
            code = []
            while index < len(lines) and not lines[index].startswith("```"):
                code.append(lines[index])
                index += 1
            index += 1
            blocks.append(("code", language, code))
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].startswith("|"):
                rows.append(lines[index])
                index += 1
            blocks.append(("table", rows))
            continue
        image_match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image_match:
            blocks.append(("image", image_match.group(1), image_match.group(2)))
            index += 1
            continue
        heading_match = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading_match:
            blocks.append(("heading", len(heading_match.group(1)), heading_match.group(2).strip()))
            index += 1
            continue
        if line.startswith(">"):
            quote = []
            while index < len(lines) and lines[index].startswith(">"):
                quote.append(re.sub(r"^>\s?", "", lines[index]))
                index += 1
            blocks.append(("quote", quote))
            continue
        if re.match(r"^\s*[-*]\s+", line):
            items = []
            while index < len(lines) and re.match(r"^\s*[-*]\s+", lines[index]):
                items.append(lines[index])
                index += 1
            blocks.append(("bullets", items))
            continue
        if re.match(r"^\s*\d+\.\s+", line):
            items = []
            while index < len(lines) and re.match(r"^\s*\d+\.\s+", lines[index]):
                items.append(lines[index])
                index += 1
            blocks.append(("numbers", items))
            continue
        if line.strip() == "---":
            blocks.append(("hr",))
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        paragraph = []
        while index < len(lines) and lines[index].strip():
            candidate = lines[index]
            if (
                candidate.startswith(("#", "|", ">", "```"))
                or re.match(r"^\s*[-*]\s+", candidate)
                or re.match(r"^\s*\d+\.\s+", candidate)
                or candidate.strip() == "---"
            ):
                break
            paragraph.append(candidate.strip())
            index += 1
        if paragraph:
            blocks.append(("paragraph", " ".join(paragraph)))
    return blocks


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(BODY)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.space_before = Pt(0)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), EAST_ASIA)

    for name, size, color, before, after in (
        ("Heading 1", 20, INK, 0, 10),
        ("Heading 2", 15, INK, 16, 7),
        ("Heading 3", 12, INK, 12, 4),
        ("Heading 4", 10.5, INK, 8, 3),
    ):
        style = styles[name]
        style.font.name = LATIN
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.2
        style_rPr = style.element.get_or_add_rPr()
        style_rFonts = style_rPr.find(qn("w:rFonts"))
        if style_rFonts is None:
            style_rFonts = OxmlElement("w:rFonts")
            style_rPr.append(style_rFonts)
        style_rFonts.set(qn("w:eastAsia"), EAST_ASIA)


def setup_page(document):
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)
    section.different_first_page_header_footer = True

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_paragraph.add_run("Workbench  ·  模型追踪")
    set_run_font(run, size=8, color=FAINT)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = footer_paragraph.add_run("2026-10-23  ·  Workbench 内部资料")
    set_run_font(run, size=8, color=FAINT)
    footer_paragraph.add_run("\t\t")
    run = footer_paragraph.add_run("第 ")
    set_run_font(run, size=8, color=FAINT)
    add_page_field(footer_paragraph)
    run = footer_paragraph.add_run(" 页")
    set_run_font(run, size=8, color=FAINT)
    tab_stops = footer_paragraph.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Cm(CONTENT_WIDTH_CM), alignment=WD_TAB_ALIGNMENT.RIGHT)


def add_cover(document):
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)
    run = spacer.add_run("")
    set_run_font(run, size=1)

    bar = document.add_table(rows=1, cols=1)
    bar.autofit = False
    bar.alignment = WD_TABLE_ALIGNMENT.CENTER
    bar.columns[0].width = Cm(CONTENT_WIDTH_CM)
    cell = bar.cell(0, 0)
    cell.width = Cm(CONTENT_WIDTH_CM)
    shade_cell(cell, INK)
    set_cell_margins(cell, top=20, bottom=20, left=0, right=0)
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    cell.paragraphs[0].paragraph_format.line_spacing = 1.0
    run = cell.paragraphs[0].add_run(" ")
    set_run_font(run, size=2)

    document.add_paragraph().paragraph_format.space_after = Pt(20)
    eyebrow = document.add_paragraph()
    eyebrow.paragraph_format.space_after = Pt(4)
    run = eyebrow.add_run("WORKBENCH  /  MODEL TRACE")
    set_run_font(run, size=9, bold=True, color=GOLD)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    title.paragraph_format.line_spacing = 1.05
    run = title.add_run("模型追踪")
    set_run_font(run, size=34, bold=True, color=INK)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(4)
    run = subtitle.add_run("操作手册")
    set_run_font(run, size=17, color=MUTED)

    deck = document.add_paragraph()
    deck.paragraph_format.space_after = Pt(14)
    run = deck.add_run("怎么看  ·  怎么判断  ·  怎么定位改动文件")
    set_run_font(run, size=10, color=FAINT)

    rule = document.add_paragraph()
    rule.paragraph_format.space_after = Pt(18)
    set_paragraph_border(rule, color=GOLD, size=12, space=1)

    lead = document.add_paragraph()
    lead.paragraph_format.line_spacing = 1.65
    lead.paragraph_format.space_after = Pt(14)
    add_inline(
        lead,
        "估值模型从第一次生成开始，每次变化追加一条记录；每条只回答三件事：变了什么、为什么变、改了哪些文件。",
        size=12,
        color=BODY,
    )

    metadata = document.add_table(rows=3, cols=2)
    metadata.autofit = False
    metadata.alignment = WD_TABLE_ALIGNMENT.LEFT
    widths = (Cm(3.2), Cm(10.5))
    rows = [
        ("版本", "2026-10-23  ·  v1.0"),
        ("适用", "Workbench 8099  ·  p-c1db281e"),
        ("读者", "投研、产品、使用者（技术说明附后）"),
    ]
    for row_index, (label, value) in enumerate(rows):
        for column_index, text in enumerate((label, value)):
            cell = metadata.cell(row_index, column_index)
            cell.width = widths[column_index]
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.2
            run = cell.paragraphs[0].add_run(text)
            set_run_font(
                run,
                size=9 if column_index else 8.5,
                bold=column_index == 0,
                color=FAINT if column_index == 0 else BODY,
            )
    set_table_borders(metadata, color=LINE, size=2)

    document.add_paragraph().paragraph_format.space_after = Pt(10)
    note = document.add_paragraph()
    note.paragraph_format.line_spacing = 1.5
    note.paragraph_format.space_before = Pt(10)
    add_inline(note, "本文是操作手册；技术说明见第 11 节。设计取舍见 ", size=9, color=MUTED)
    add_hyperlink(note, "model-trace.md", (DOCS / "model-trace.md").as_uri(), size=9)
    add_inline(note, "。截图来自真实项目，判断卡为项目级入口。", size=9, color=MUTED)

    bottom_rule = document.add_paragraph()
    bottom_rule.paragraph_format.space_before = Pt(188)
    bottom_rule.paragraph_format.space_after = Pt(5)
    set_paragraph_border(bottom_rule, color=LINE, size=4, space=1)

    bottom_note = document.add_paragraph()
    bottom_note.paragraph_format.space_after = Pt(0)
    run = bottom_note.add_run("从生成到每一次中途改变")
    set_run_font(run, size=9, color=FAINT)
    run = bottom_note.add_run("                                                                2026 · Workbench")
    set_run_font(run, size=9, color=FAINT)

    document.add_page_break()


def add_toc(document, blocks):
    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(12)
    run = heading.add_run("目录")
    set_run_font(run, size=20, bold=True, color=INK)
    set_paragraph_border(heading, color=LINE, size=6, space=4)

    sections = [(level, text) for kind, level, text in ((block[0], block[1], block[2]) for block in blocks if block[0] == "heading") if level == 2]
    for number, (_, text) in enumerate(sections, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(7)
        paragraph.paragraph_format.line_spacing = 1.25
        prefix = paragraph.add_run(f"{number:02d}   ")
        set_run_font(prefix, size=9, bold=True, color=GOLD)
        title_run = paragraph.add_run(re.sub(r"^\d+\.\s*", "", clean_md(text)))
        set_run_font(title_run, size=11, color=BODY)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(16)
    note.paragraph_format.line_spacing = 1.5
    add_inline(
        note,
        "说明：截图中的项目为长鑫科技（p-c1db281e）。#4 是归母口径提案，当前目标价仍是 ¥92；提案只有在人工采纳并写入主报告后才会生效。",
        size=9,
        color=MUTED,
    )
    document.add_page_break()


def add_code_block(document, lines):
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Cm(CONTENT_WIDTH_CM)
    cell = table.cell(0, 0)
    cell.width = Cm(CONTENT_WIDTH_CM)
    shade_cell(cell, CODE_BG)
    set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
    set_table_borders(table, color="E4DED2", size=2)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.12
    if not lines:
        lines = [""]
    for index, line in enumerate(lines):
        if index:
            paragraph = cell.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.12
        run = paragraph.add_run(line if line else " ")
        set_run_font(run, name=MONO, size=8, color="3A372F")
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_quote(document, lines):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.45)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.45
    set_paragraph_border(paragraph, edge="left", color=GOLD, size=12, space=8)
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        add_inline(paragraph, line, size=9.5, color=MUTED, italic=False)


def add_bullets(document, items):
    for item in items:
        match = re.match(r"^(\s*)[-*]\s+(.*)$", item)
        if not match:
            continue
        indent = len(match.group(1).replace("\t", "  ")) // 2
        text = match.group(2).strip()
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.55 + 0.4 * indent)
        paragraph.paragraph_format.first_line_indent = Cm(-0.35)
        paragraph.paragraph_format.space_after = Pt(2.5)
        paragraph.paragraph_format.line_spacing = 1.4
        prefix = paragraph.add_run("–  " if indent else "•  ")
        set_run_font(prefix, size=10, bold=True, color=GOLD)
        add_inline(paragraph, text, size=10.5, color=BODY)


def add_numbers(document, items):
    for number, item in enumerate(items, start=1):
        match = re.match(r"^\s*\d+\.\s+(.*)$", item)
        if not match:
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.62)
        paragraph.paragraph_format.first_line_indent = Cm(-0.62)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.4
        tab_stops = paragraph.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Cm(0.62))
        prefix = paragraph.add_run(f"{number:02d}\t")
        set_run_font(prefix, size=9.5, bold=True, color=GOLD)
        add_inline(paragraph, match.group(1).strip(), size=10.5, color=BODY)


def add_table(document, raw_rows):
    rows = []
    for raw in raw_rows:
        cells = [cell.strip() for cell in raw.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells if cell):
            continue
        rows.append(cells)
    if not rows:
        return
    columns = max(len(row) for row in rows)
    rows = [row + [""] * (columns - len(row)) for row in rows]
    table = document.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_fixed(table)
    set_table_borders(table, color=LINE, size=4)

    weights = []
    for column in range(columns):
        lengths = [len(clean_md(row[column])) for row in rows]
        weights.append(max(8, min(34, max(lengths) if lengths else 8)))
    total = sum(weights)
    widths = [CONTENT_WIDTH_CM * weight / total for weight in weights]
    for column, width in enumerate(widths):
        table.columns[column].width = Cm(width)

    for row_index, row in enumerate(rows):
        table_row = table.rows[row_index]
        set_row_cant_split(table_row)
        for column, value in enumerate(row):
            cell = table_row.cells[column]
            cell.width = Cm(widths[column])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=55, bottom=55, left=80, right=80)
            if row_index == 0:
                shade_cell(cell, INK)
            elif row_index % 2 == 0:
                shade_cell(cell, CARD)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.25
            paragraph.paragraph_format.keep_with_next = row_index < len(rows) - 1
            add_inline(
                paragraph,
                value,
                size=8.5,
                bold=row_index == 0,
                color="FFFFFF" if row_index == 0 else BODY,
            )
        if row_index == 0:
            set_repeat_header(table_row)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(document, alt, relative_path, figure_number):
    path = (DOCS / relative_path).resolve()
    if not path.exists():
        paragraph = document.add_paragraph()
        add_inline(paragraph, f"[缺图：{alt}]", size=9, color=MUTED)
        return figure_number
    with Image.open(path) as image:
        width, height = image.size
    aspect = width / height if height else 1
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.space_before = Pt(6)
    picture_paragraph.paragraph_format.space_after = Pt(2)
    picture_paragraph.paragraph_format.keep_with_next = True
    run = picture_paragraph.add_run()
    if aspect > 1.25:
        run.add_picture(str(path), width=Cm(16.2))
    elif aspect >= 0.85:
        run.add_picture(str(path), width=Cm(12))
    else:
        run.add_picture(str(path), height=Cm(12))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False
    run = caption.add_run(f"图 {figure_number}  {alt}")
    set_run_font(run, size=8.5, color=FAINT)
    return figure_number + 1


def add_body(document, blocks):
    figure = 1
    skipped_title = False
    for index, block in enumerate(blocks):
        kind = block[0]
        next_kind = blocks[index + 1][0] if index + 1 < len(blocks) else None
        if kind == "heading":
            level = block[1]
            text = block[2]
            if level == 1 and not skipped_title:
                skipped_title = True
                continue
            if level == 2:
                number_match = re.match(r"^(\d+)\.", text)
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(18)
                paragraph.paragraph_format.space_after = Pt(8)
                paragraph.paragraph_format.keep_with_next = True
                number = number_match.group(1) if number_match else ""
                if number:
                    prefix = paragraph.add_run(f"{number}  ")
                    set_run_font(prefix, size=10, bold=True, color=GOLD)
                run = paragraph.add_run(re.sub(r"^\d+\.\s*", "", text))
                set_run_font(run, size=15, bold=True, color=INK)
                set_paragraph_border(paragraph, color=LINE, size=6, space=4)
            elif level == 3:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(12)
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.keep_with_next = True
                run = paragraph.add_run(clean_md(text))
                set_run_font(run, size=12, bold=True, color=INK)
            elif level == 4:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(8)
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.keep_with_next = True
                run = paragraph.add_run(clean_md(text))
                set_run_font(run, size=10.5, bold=True, color=MUTED)
        elif kind == "paragraph":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.line_spacing = 1.45
            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.paragraph_format.keep_together = True
            if next_kind in ("image", "table", "code"):
                paragraph.paragraph_format.keep_with_next = True
            add_inline(paragraph, block[1], size=10.5, color=BODY)
        elif kind == "quote":
            add_quote(document, block[1])
        elif kind == "bullets":
            add_bullets(document, block[1])
        elif kind == "numbers":
            add_numbers(document, block[1])
        elif kind == "code":
            add_code_block(document, block[2])
        elif kind == "table":
            add_table(document, block[1])
        elif kind == "image":
            figure = add_image(document, block[1], block[2], figure)
        elif kind == "hr":
            continue


def main():
    source = SOURCE.read_text(encoding="utf-8")
    blocks = parse_blocks(source)

    document = Document()
    configure_styles(document)
    setup_page(document)
    document.core_properties.title = "模型追踪：操作手册"
    document.core_properties.subject = "Workbench 模型追踪操作手册（附技术说明）"
    document.core_properties.author = "Workbench"
    document.core_properties.comments = "由 docs/model-trace-guide.md 生成"

    add_cover(document)
    add_toc(document, blocks)
    add_body(document, blocks)

    document.save(OUTPUT)
    print(f"saved {OUTPUT}")

    reopened = Document(OUTPUT)
    print("paragraphs:", len(reopened.paragraphs))
    print("tables:", len(reopened.tables))
    print("images:", len(reopened.inline_shapes))
    print("sections:", len(reopened.sections))


if __name__ == "__main__":
    main()
